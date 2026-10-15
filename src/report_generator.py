
import json
from pathlib import Path

import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


PROJECT_ROOT = Path(__file__).resolve().parents[1]
OUTPUT_FOLDER = PROJECT_ROOT / "output"


def generate_excel_report(
    json_report_path=None,
    excel_report_path=None,
):
    """Generate a formatted Excel quality-check report."""

    if json_report_path is None:
        json_report_path = (
            OUTPUT_FOLDER / "quality_check_results.json"
        )

    if excel_report_path is None:
        excel_report_path = (
            OUTPUT_FOLDER / "quality_check_report.xlsx"
        )

    json_report_path = Path(json_report_path)
    excel_report_path = Path(excel_report_path)

    if not json_report_path.exists():
        raise FileNotFoundError(
            f"JSON report not found: {json_report_path}"
        )

    with json_report_path.open("r", encoding="utf-8") as file:
        report = json.load(file)

    findings = report.get("findings", [])

    # Convert nested values into readable Excel cells.
    findings_for_excel = []

    for finding in findings:
        row = finding.copy()

        for key, value in row.items():
            if isinstance(value, (dict, list)):
                row[key] = json.dumps(
                    value,
                    ensure_ascii=False,
                    default=str,
                )

        findings_for_excel.append(row)

    findings_df = pd.DataFrame(findings_for_excel)

    # Overall summary.
    summary_df = pd.DataFrame([
        {
            "Metric": "Datasets loaded",
            "Value": report.get("total_datasets_loaded", 0),
        },
        {
            "Metric": "Files with reading errors",
            "Value": report.get("total_read_errors", 0),
        },
        {
            "Metric": "Total findings",
            "Value": len(findings),
        },
    ])

    # Findings grouped by rule and dataset.
    if findings:
        raw_findings_df = pd.DataFrame(findings)

        rule_summary = (
            raw_findings_df.groupby("rule_id")
            .size()
            .reset_index(name="finding_count")
        )

        dataset_summary = (
            raw_findings_df.groupby("dataset")
            .size()
            .reset_index(name="finding_count")
            .sort_values("finding_count", ascending=False)
        )
    else:
        rule_summary = pd.DataFrame(
            columns=["rule_id", "finding_count"]
        )
        dataset_summary = pd.DataFrame(
            columns=["dataset", "finding_count"]
        )

    OUTPUT_FOLDER.mkdir(parents=True, exist_ok=True)
    excel_report_path.parent.mkdir(
        parents=True, exist_ok=True
    )

    with pd.ExcelWriter(
        excel_report_path,
        engine="openpyxl",
    ) as writer:
        summary_df.to_excel(
            writer, sheet_name="Summary", index=False
        )
        rule_summary.to_excel(
            writer, sheet_name="Findings by Rule", index=False
        )
        dataset_summary.to_excel(
            writer, sheet_name="Findings by Dataset", index=False
        )
        findings_df.to_excel(
            writer, sheet_name="Findings", index=False
        )

        # Format every worksheet.
        for worksheet in writer.sheets.values():
            worksheet.freeze_panes = "A2"
            worksheet.auto_filter.ref = worksheet.dimensions

            # Format the header row.
            for cell in worksheet[1]:
                cell.font = Font(
                    bold=True,
                    color="FFFFFF",
                )
                cell.fill = PatternFill(
                    fill_type="solid",
                    fgColor="1F4E78",
                )
                cell.alignment = Alignment(
                    horizontal="center",
                    vertical="center",
                    wrap_text=True,
                )

            worksheet.row_dimensions[1].height = 30

            # Wrap long text and calculate useful column widths.
            for column_cells in worksheet.columns:
                column_letter = get_column_letter(
                    column_cells[0].column
                )
                header = str(column_cells[0].value or "")
                max_length = len(header)

                for cell in column_cells[1:]:
                    value = cell.value

                    if value is not None:
                        text_value = str(value)
                        max_length = max(
                            max_length,
                            min(
                                max(
                                    len(line)
                                    for line in text_value.splitlines()
                                ),
                                55,
                            ),
                        )

                    cell.alignment = Alignment(
                        vertical="top",
                        wrap_text=True,
                    )

                worksheet.column_dimensions[
                    column_letter
                ].width = min(max(max_length + 2, 12), 40)

            # Give long descriptions and JSON values more room.
            for row_number in range(2, worksheet.max_row + 1):
                max_lines = 1

                for cell in worksheet[row_number]:
                    if cell.value is not None:
                        lines = str(cell.value).count("\n") + 1
                        max_lines = max(max_lines, lines)

                worksheet.row_dimensions[row_number].height = min(
                    max(20, max_lines * 18),
                    90,
                )

    print("Excel report generated successfully.")
    print("Total findings:", len(findings))
    print("Report saved to:", excel_report_path)

    return excel_report_path
